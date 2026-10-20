from pathlib import Path
from collections import Counter
from datetime import datetime
import hashlib
import re

MAX_FILES_PER_SCAN = 5000
MAX_SUMMARY_FILE_BYTES = 20 * 1024 * 1024  # 20 MiB

SAFE_CATEGORIES = {
    "Images": {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg", ".tif", ".tiff", ".heic"},
    "PDFs": {".pdf"},
    "Spreadsheets": {".xlsx", ".xls", ".xlsm", ".ods", ".csv", ".tsv"},
    "Documents": {".doc", ".docx", ".odt", ".rtf", ".txt", ".md", ".pages"},
    "Presentations": {".ppt", ".pptx", ".odp", ".key"},
    "Videos": {".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv"},
    "Audio": {".mp3", ".wav", ".aac", ".flac", ".ogg", ".m4a"},
    "Archives": {".zip", ".rar", ".7z", ".tar", ".gz"},
    "Code": {".py", ".js", ".ts", ".html", ".css", ".java", ".c", ".cpp", ".json", ".yml", ".yaml", ".sql"},
}

def classify_file(path: Path) -> str:
    suffix = path.suffix.lower()
    for category, extensions in SAFE_CATEGORIES.items():
        if suffix in extensions:
            return category
    return "Other"

def format_bytes(value: int) -> str:
    value = float(value)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if value < 1024 or unit == "TB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} TB"

def scan_folder(root: Path, recursive: bool = True):
    """Scan regular files without following symlinks; cap work to limit resource abuse."""
    results, errors = [], []
    try:
        root = root.expanduser().resolve(strict=True)
        if not root.is_dir():
            return [], [str(root)]
        # os.walk does not follow directory symlinks by default.
        import os
        for current, dirs, files in os.walk(root, topdown=True, followlinks=False):
            current_path = Path(current)
            # Explicitly prune symlinked directories and hidden VCS/virtual-env folders.
            dirs[:] = [d for d in dirs if not (current_path / d).is_symlink() and d not in {".git", ".venv", "venv", "node_modules", "__pycache__"}]
            if not recursive and current_path != root:
                dirs[:] = []
                continue
            for filename in files:
                path = current_path / filename
                try:
                    if path.is_symlink() or not path.is_file():
                        continue
                    resolved = path.resolve(strict=True)
                    if not resolved.is_relative_to(root):
                        errors.append("Skipped a file outside the selected folder.")
                        continue
                    stat = resolved.stat()
                    results.append({
                        "name": resolved.name,
                        "path": str(resolved),
                        "size": stat.st_size,
                        "modified": datetime.fromtimestamp(stat.st_mtime).strftime("%d %b %Y, %I:%M %p"),
                        "category": classify_file(resolved),
                    })
                    if len(results) >= MAX_FILES_PER_SCAN:
                        errors.append(f"Scan stopped at the safety limit of {MAX_FILES_PER_SCAN:,} files.")
                        return results, errors
                except (OSError, PermissionError, RuntimeError):
                    errors.append("One or more inaccessible files were skipped.")
    except (OSError, PermissionError, RuntimeError):
        errors.append("The selected folder could not be accessed safely.")
    return results, errors

def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    """Hash a regular non-symlink file in bounded chunks."""
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise OSError("Refusing to hash a symlink or non-regular file.")
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()

def _read_text(path: Path, limit: int = 12000) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="ignore")[:limit]
    except OSError:
        return ""

def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        chunks = []
        for page in reader.pages[:8]:
            text = page.extract_text() or ""
            if text:
                chunks.append(text)
            if sum(len(x) for x in chunks) > 12000:
                break
        return "\n".join(chunks)[:12000]
    except Exception:
        return ""

def _excel_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        chunks = []
        for sheet in workbook.worksheets[:5]:
            chunks.append(f"Sheet: {sheet.title}")
            rows_seen = 0
            for row in sheet.iter_rows(values_only=True):
                values = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if values:
                    chunks.append(" | ".join(values[:10]))
                    rows_seen += 1
                if rows_seen >= 30 or sum(len(x) for x in chunks) > 10000:
                    break
        workbook.close()
        return "\n".join(chunks)[:12000]
    except Exception:
        return ""

def _docx_text(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables[:3]:
            for row in table.rows[:20]:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(chunks)[:12000]
    except Exception:
        return ""

def _csv_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8-sig", errors="ignore")[:12000]
    except OSError:
        return ""

def _words(text: str):
    return re.findall(r"[A-Za-z0-9][A-Za-z0-9’'_-]*", text.lower())

def _rule_summary(text: str, path: Path) -> str:
    """Small offline summary: pick informative lines/phrases, then enforce <=15 words."""
    name = path.stem.replace("_", " ").replace("-", " ").strip()
    cleaned_lines = []
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line).strip(" \t|,;:-")
        if len(line) >= 3:
            cleaned_lines.append(line)
    # Prefer lines that resemble headings or contain several meaningful words.
    stop = {"the","and","for","with","from","this","that","are","was","were","into","your","have","has","not","but","all","sheet"}
    candidates = []
    for line in cleaned_lines:
        tokens = _words(line)
        if not tokens:
            continue
        score = min(len(tokens), 18) + min(len(line), 100) / 100
        score += sum(1 for token in tokens if token not in stop) * 0.15
        candidates.append((score, line))
    if candidates:
        candidates.sort(key=lambda item: item[0], reverse=True)
        chosen = candidates[0][1]
        # Avoid an excessively long table row becoming the entire summary.
        chosen = re.split(r"[.!?]\s+", chosen)[0]
        tokens = re.findall(r"\S+", chosen)
        summary = " ".join(tokens[:15]).strip(" ,;:-")
        if len(tokens) > 15:
            summary = summary.rstrip(".,;:") + "…"
        if summary:
            return summary
    if name:
        fallback = f"{name} file. Open it to review its contents."
        return " ".join(fallback.split()[:15])
    return "No readable text found. This may be scanned, protected, or unsupported."

def get_file_summary(path: Path) -> str:
    path = Path(path)
    try:
        if path.is_symlink() or not path.is_file():
            return "Preview unavailable for links or non-regular files."
        if path.stat().st_size > MAX_SUMMARY_FILE_BYTES:
            return "Preview skipped: file exceeds the 20 MiB safety limit."
    except OSError:
        return "Preview unavailable because the file could not be accessed."
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _pdf_text(path)
    elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        text = _excel_text(path)
    elif suffix == ".docx":
        text = _docx_text(path)
    elif suffix in {".txt", ".md"}:
        text = _read_text(path)
    elif suffix == ".csv":
        text = _csv_text(path)
    else:
        return "Preview summary is not available for this file type."
    if not text.strip():
        if suffix == ".pdf":
            return "No readable PDF text found; it may be scanned, encrypted, or image-only."
        return "No readable text found. The document may be empty or use unsupported formatting."
    return _rule_summary(text, path)
